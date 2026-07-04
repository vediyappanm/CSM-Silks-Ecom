from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "CSM-Silks-Production-Documentation.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(32, 32, 32)
MUTED = RGBColor(90, 90, 90)
HEADER_FILL = "E8EEF5"
LIGHT_FILL = "F4F6F9"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths_dxa: list[int]) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl.insert(1, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, width in enumerate(widths_dxa):
            if idx >= len(row.cells):
                continue
            cell = row.cells[idx]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def mark_header_row(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


def add_run(paragraph, text: str, bold=False, italic=False, color=None, size=None):
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color
    if size:
        run.font.size = Pt(size)
    return run


def add_para(doc, text: str = "", style: str | None = None):
    p = doc.add_paragraph(style=style)
    if text:
        p.add_run(text)
    return p


def add_bullets(doc, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_numbers(doc, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)


def add_code_block(doc, lines: list[str]) -> None:
    for line in lines:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(line)
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(40, 40, 40)


def add_callout(doc, title: str, body: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    set_table_width(table, [9360])
    mark_header_row(table.rows[0])
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_FILL)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    add_run(p, title + ": ", bold=True, color=DARK_BLUE)
    add_run(p, body)
    doc.add_paragraph()


def add_table(doc, headers: list[str], rows: list[list[str]], widths_dxa: list[int]):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    set_table_width(table, widths_dxa)
    for i, header in enumerate(headers):
        cell = table.cell(0, i)
        set_cell_shading(cell, HEADER_FILL)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        add_run(p, header, bold=True, color=DARK_BLUE)
    mark_header_row(table.rows[0])
    for row_data in rows:
        row = table.add_row()
        for i, value in enumerate(row_data):
            cell = row.cells[i]
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.add_run(value)
    set_table_width(table, widths_dxa)
    doc.add_paragraph()
    return table


def configure_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ["List Bullet", "List Number"]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25


def add_header_footer(doc: Document) -> None:
    section = doc.sections[0]
    header_p = section.header.paragraphs[0]
    header_p.text = "CSM Silks Production Documentation"
    header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_p.runs[0].font.size = Pt(9)
    header_p.runs[0].font.color.rgb = MUTED

    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_p.add_run("Production handover manual")
    footer_p.runs[0].font.size = Pt(9)
    footer_p.runs[0].font.color.rgb = MUTED


def build_document() -> None:
    doc = Document()
    configure_styles(doc)
    add_header_footer(doc)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(3)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_run(title, "CSM Silks Ecommerce Platform", size=24, color=DARK_BLUE, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(12)
    add_run(subtitle, "End-to-End Production Documentation and Operations Manual", size=14, color=MUTED)

    add_table(
        doc,
        ["Document field", "Value"],
        [
            ["Project", "CSM Silks single-brand textile ecommerce platform"],
            ["Repository", "C:\\Users\\VediyappanMFinspot\\Music\\CSM-Silks-Ecom"],
            ["Generated", date.today().isoformat()],
            ["Primary audience", "Developers, admin operators, QA, deployment owners, and business stakeholders"],
            ["Runtime stack", "Django 6 / DRF / Channels / Celery / Redis / PostgreSQL / React 19 / Vite 8"],
        ],
        [2400, 6960],
    )

    add_callout(
        doc,
        "Production readiness position",
        "The application is a real API-backed ecommerce system with customer OTP auth, admin operations, live catalog/cart/order flows, Razorpay/COD, shipments, returns, loyalty, notifications, WebSocket updates, and production deployment checks. Production launch still requires live provider credentials, domain/DNS, SSL, PostgreSQL, Redis, backup policy, and final payment/courier webhook registration.",
    )

    doc.add_heading("Table of Contents", level=1)
    add_numbers(
        doc,
        [
            "Executive summary",
            "Business scope and user roles",
            "Architecture and repository structure",
            "Local setup and developer workflow",
            "Environment variables and production configuration",
            "Data model overview",
            "Customer journey end to end",
            "Admin operations end to end",
            "API contract",
            "Realtime and background jobs",
            "Payments, shipping, invoices, returns, loyalty, and notifications",
            "Security and compliance controls",
            "Deployment and CI/CD",
            "Operational runbooks",
            "Testing, QA, and acceptance checklist",
            "Known limits and next roadmap",
        ],
    )

    doc.add_heading("1. Executive Summary", level=1)
    add_para(
        doc,
        "CSM Silks is a single-brand textile ecommerce platform built for a retailer that owns its own inventory. It follows an Amazon/Flipkart-like customer flow while keeping the business domain specific to sarees, dhotis, silk shirts, variants, blouse options, stock, GST invoices, order tracking, returns, rewards, and assisted customer communications.",
    )
    add_para(
        doc,
        "The backend is Django/DRF with a practical MVVM split: Django ORM models represent domain state, serializers act as view models/presenters for React screens, APIViews/ViewSets remain thin HTTP surfaces, and service modules hold checkout, pricing, inventory, payment, shipping, invoice, notification, and AI logic. The frontend is a React/Vite SPA with authenticated customer routes, admin console, shared API client, theme provider, and realtime WebSocket helpers.",
    )
    add_table(
        doc,
        ["Layer", "Production responsibility", "Current implementation"],
        [
            ["Storefront", "Customer browsing, product detail, cart, checkout, orders, tracking, wishlist, notifications, try-on", "React 19 SPA in frontend/src/pages with API-backed state"],
            ["Admin", "Catalog, variants, images, stock, orders, shipments, returns, coupons, reports, audit logs", "Admin console in frontend/src/pages/Admin.tsx plus AdminCatalogManager"],
            ["API", "Auth, commerce, inventory, orders, payments, shipping, analytics, AI", "Django apps under backend/ with /api prefix"],
            ["Realtime", "Live catalog/inventory/order/notification updates", "Django Channels endpoints under /ws and React realtime client"],
            ["Async jobs", "Reservation expiry, notifications, reports, email/PDF work", "Celery worker and beat backed by Redis"],
            ["Production infra", "Django ASGI, frontend static serving, Nginx, PostgreSQL, Redis", "docker-compose.prod.yml and nginx.conf"],
        ],
        [1500, 3600, 4260],
    )

    doc.add_heading("2. Business Scope and User Roles", level=1)
    doc.add_heading("2.1 In Scope for V1", level=2)
    add_bullets(
        doc,
        [
            "Single-brand catalog for CSM-owned textile inventory.",
            "Textile-specific product data: fabric, color, occasion, zari, blouse option, size, care, SKU, image gallery, and stock per variant.",
            "Customer account access through OTP and admin access through email/password.",
            "Cart, wishlist, coupons, checkout summary, COD and Razorpay payment flows.",
            "Order lifecycle from draft/confirmed through packing, shipping, delivery, RTO, cancellation, return, refund, and invoice.",
            "Admin operations for product upload, collection creation, stock adjustment, order workflow, shipment labels, return decisions, coupons, customers, reports, and audit logs.",
            "Customer notifications through application feed plus optional email, SMS, and WhatsApp provider configuration.",
        ],
    )
    doc.add_heading("2.2 Roles", level=2)
    add_table(
        doc,
        ["Role", "Authentication", "Capabilities"],
        [
            ["Customer", "Phone OTP; optional email for OTP/profile", "Shop, cart, wishlist, checkout, track orders, download invoice, request return, review delivered products, manage loyalty."],
            ["Admin / super admin", "Email and password; seeded admin available locally", "Manage catalog, inventory, order workflow, shipments, returns, coupons, analytics, audit logs."],
            ["Anonymous visitor", "No session", "Redirected to login/signup before protected storefront routes; admin route shows admin login."],
            ["System worker", "Celery process", "Releases expired reservations, sends notifications, runs scheduled background work."],
        ],
        [1900, 2200, 5260],
    )

    doc.add_heading("3. Architecture and Repository Structure", level=1)
    add_code_block(
        doc,
        [
            "backend/",
            "  csm_backend/      settings, urls, ASGI, WebSocket routing, Celery, health/readiness",
            "  accounts/         customer/admin auth, OTP, profile, addresses, production checks",
            "  catalog/          categories, collections, products, variants, images, facets, realtime",
            "  inventory/        stock ledger, reservations, unsold alerts, reservation expiry task",
            "  cart/             cart, cart items, coupon preview, wishlist",
            "  orders/           checkout, order lifecycle, invoices, returns, coupons, realtime",
            "  payments/         Razorpay order/verify/webhook/refund and COD support",
            "  shipping/         shipments, events, labels, manifests, Shiprocket/manual adapter",
            "  loyalty/          points ledger and rewards",
            "  notifications/    notification logs, Resend/Twilio/Gupshup services, realtime feed",
            "  analytics/        admin KPIs, reports, customers, audit logs",
            "  reviews/          verified product reviews",
            "  ai/               Claude-assisted try-on, voice search, recommendations",
            "frontend/",
            "  src/pages/        customer and admin route screens",
            "  src/features/     feature-level UI such as admin catalog manager and catalog cards",
            "  src/lib/          API client, realtime client, order lifecycle helpers",
            "  src/store/        app state and theme providers",
            "  src/styles/       tokens, base, layout, product, commerce, admin CSS",
        ],
    )
    add_table(
        doc,
        ["Architectural concern", "Design decision"],
        [
            ["Backend pattern", "Practical MVVM: models = Django ORM, view models = serializers, views = thin DRF APIViews, services/selectors = business and query logic."],
            ["Frontend pattern", "Route-level screens consume a single API client and AppContext; common components handle navigation, product visuals, skeletons, and error boundaries."],
            ["Database", "SQLite is allowed for local dev; PostgreSQL is required for production and enforced by deploy checks."],
            ["Realtime", "Django Channels with JWT subprotocol auth; Redis channel layer required in production."],
            ["Background jobs", "Celery worker and beat use Redis broker/result backend; beat releases expired stock reservations."],
            ["Infrastructure", "Daphne ASGI app behind Nginx; Vite frontend built and served as static app in production compose."],
        ],
        [2300, 7060],
    )

    doc.add_heading("4. Local Setup and Developer Workflow", level=1)
    doc.add_heading("4.1 Native Local Setup", level=2)
    add_numbers(
        doc,
        [
            "Create .env from .env.example and keep .env untracked.",
            "Install Python dependencies with pip install -r requirements.txt.",
            "Run python backend/manage.py migrate.",
            "Run python backend/manage.py seed_csm to seed products, admin, customer, and demo commerce data.",
            "Start the backend ASGI app from backend: daphne -b 0.0.0.0 -p 8000 csm_backend.asgi:application.",
            "Install frontend dependencies from frontend: npm install.",
            "Start Vite: npm run dev. It proxies /api and /ws to http://127.0.0.1:8000.",
        ],
    )
    add_code_block(
        doc,
        [
            "python backend/manage.py migrate",
            "python backend/manage.py seed_csm",
            "cd backend",
            "daphne -b 0.0.0.0 -p 8000 csm_backend.asgi:application",
            "cd ../frontend",
            "npm install",
            "npm run dev",
        ],
    )
    doc.add_heading("4.2 Local Accounts", level=2)
    add_table(
        doc,
        ["Account", "Value"],
        [
            ["Admin", "admin@csmsilks.com / admin123"],
            ["Customer OTP phone", "+918888888888"],
            ["Dev OTP fallback", "Only returned when DEBUG=True and OTP_DEV_FALLBACK_ENABLED=True"],
            ["Admin URL", "http://localhost:5173/admin"],
            ["Customer URL", "http://localhost:5173/login or http://localhost:5173/signup"],
        ],
        [2600, 6760],
    )
    doc.add_heading("4.3 Docker Development", level=2)
    add_para(doc, "The development Docker Compose stack includes API, Celery worker, Celery beat, PostgreSQL, Redis, and optional pgAdmin.")
    add_code_block(doc, ["docker compose up", "docker compose --profile dev up"])

    doc.add_heading("5. Environment Variables and Production Configuration", level=1)
    add_para(
        doc,
        "The application is intentionally strict about production configuration. With DEBUG=False or APP_ENV=production, deploy checks block unsafe defaults such as SQLite, in-memory realtime, missing Redis/Celery settings, missing Razorpay credentials/webhook secret, no live OTP channel, no notification channel, and Shiprocket without credentials when selected.",
    )
    add_table(
        doc,
        ["Variable group", "Important variables", "Production guidance"],
        [
            ["Core", "APP_ENV, DEBUG, SECRET_KEY, ALLOWED_HOSTS, CORS_ALLOWED_ORIGINS, CSRF_TRUSTED_ORIGINS", "Use APP_ENV=production, DEBUG=False, long random SECRET_KEY, exact domain hosts and origins."],
            ["Database", "DATABASE_URL", "Use postgres:// or postgresql:// in production. SQLite is development only."],
            ["Redis/Celery/Channels", "REDIS_URL, CELERY_BROKER_URL, CELERY_RESULT_BACKEND, CHANNEL_LAYER_BACKEND", "Use Redis URLs and CHANNEL_LAYER_BACKEND=redis."],
            ["Security", "SECURE_SSL_REDIRECT, SESSION_COOKIE_SECURE, CSRF_COOKIE_SECURE, SECURE_HSTS_*", "Production defaults enable HTTPS redirect, secure cookies, HSTS, and forwarded proto support unless overridden."],
            ["Payments", "RAZORPAY_KEY_ID, RAZORPAY_KEY_SECRET, RAZORPAY_WEBHOOK_SECRET, PAYMENT_DEV_FALLBACK_ENABLED", "Use live/test Razorpay project keys and webhook secret; keep dev fallback false."],
            ["Shipping", "DEFAULT_COURIER_PROVIDER, SHIPROCKET_EMAIL, SHIPROCKET_PASSWORD, SHIPROCKET_WEBHOOK_SECRET", "Manual courier works without Shiprocket. If DEFAULT_COURIER_PROVIDER=shiprocket, credentials and webhook secret are required."],
            ["Notifications", "RESEND_API_KEY, RESEND_FROM_EMAIL, NOTIFICATION_EMAIL_ENABLED, WHATSAPP_ENABLED, GUPSHUP_*", "Configure at least one production customer notification channel."],
            ["OTP", "SMS_OTP_ENABLED, TWILIO_*, OTP_EMAIL_ENABLED, OTP_DEV_FALLBACK_ENABLED", "Configure Twilio SMS OTP or Resend email OTP. Disable dev fallback."],
            ["AI", "ANTHROPIC_API_KEY, ANTHROPIC_MODEL", "Required only for real Claude-backed AI try-on behavior."],
        ],
        [1700, 3600, 4060],
    )
    add_callout(
        doc,
        "Credential hygiene",
        "Never commit .env, SQLite databases, logs, screenshots, media uploads, node_modules, or build output. If any API key was ever shared in chat or committed history, rotate it before production.",
    )

    doc.add_heading("6. Data Model Overview", level=1)
    add_table(
        doc,
        ["App", "Models", "Purpose"],
        [
            ["accounts", "User, OTPChallenge, Address", "Customer/admin identity, OTP lifecycle, saved addresses."],
            ["catalog", "Category, Collection, Product, ProductVariant, ProductImage", "Textile catalog, SKU variants, images, selling metadata."],
            ["inventory", "StockLedger, StockReservation, UnsoldAlert", "Auditable stock movement, checkout stock holds, slow-moving stock alerts."],
            ["cart", "Cart, CartItem, WishlistItem", "Authenticated cart and saved products."],
            ["orders", "Coupon, Order, OrderItem, ReturnRequest", "Checkout, order lifecycle, coupon rules, item snapshots, return requests."],
            ["payments", "Payment, RazorpayWebhookEvent", "Payment attempts, captured/refunded state, webhook idempotency."],
            ["shipping", "Shipment, ShipmentEvent", "Courier state, AWB, tracking URL, labels, manifests, event timeline."],
            ["loyalty", "LoyaltyTransaction, LoyaltyReward", "Points ledger and redeemable rewards."],
            ["reviews", "ProductReview", "Verified post-delivery product feedback."],
            ["notifications", "Notification, MessageTemplate", "App feed plus email/SMS/WhatsApp send status."],
            ["analytics", "DailyReport, AdminAuditLog", "Dashboards, reports, traceability of admin changes."],
            ["ai", "TryOnSession", "AI try-on/styling request history and model metadata."],
        ],
        [1600, 3300, 4460],
    )
    add_para(
        doc,
        "ProductVariant is the sellable unit for stock-sensitive operations. Checkout, cart, inventory reservations, order items, and admin stock adjustment should always target variants, not only products.",
    )

    doc.add_heading("7. Customer Journey End to End", level=1)
    add_code_block(
        doc,
        [
            "Customer opens /login or /signup",
            "  -> POST /api/auth/otp/send",
            "  -> OTP delivered by Twilio SMS or Resend email in production",
            "  -> POST /api/auth/otp/verify",
            "  -> JWT access/refresh tokens stored by frontend client",
            "  -> Customer browses products/search/facets",
            "  -> Adds SKU variant to cart",
            "  -> Checkout validates address, cart totals, coupon, loyalty, and stock",
            "  -> Order created as COD confirmed or Razorpay payment pending/confirmed",
            "  -> Payment webhook/verification confirms prepaid order",
            "  -> Invoice, loyalty, notification, shipment/tracking lifecycle starts",
        ],
    )
    add_table(
        doc,
        ["Screen", "Route", "Key backend calls"],
        [
            ["Login/signup", "/login, /signup", "POST /api/auth/otp/send, POST /api/auth/otp/verify"],
            ["Home/catalog", "/, /womens, /mens, /search", "GET /api/products, /api/search, /api/catalog/facets, /api/categories, /api/collections"],
            ["Product detail", "/product/:gender/:id", "GET /api/products/{slug}, /api/products/{slug}/delivery, /api/products/{slug}/reviews"],
            ["Cart", "/cart", "GET/POST/DELETE /api/cart, PATCH/DELETE /api/cart/items/{id}, POST /api/cart/coupon"],
            ["Checkout", "/checkout", "GET /api/checkout/summary, POST /api/addresses, POST /api/orders, payment APIs"],
            ["Orders", "/orders", "GET /api/orders, GET /api/orders/{id}, POST /api/orders/{id}/cancel, invoice, return, review APIs"],
            ["Tracking", "/tracking, /tracking/:orderId", "GET /api/orders/track or authenticated order detail plus /ws/orders/{id}/"],
            ["Notifications", "/notifications", "GET /api/notifications, GET /api/notifications/count, PATCH /api/notifications, /ws/notifications/"],
            ["Account", "/account", "GET/PATCH /api/auth/me, addresses, loyalty, rewards, wishlist"],
            ["Try-on", "/tryon", "POST /api/ai/tryon, GET/POST /api/ai/recommend"],
        ],
        [1700, 1900, 5760],
    )

    doc.add_heading("8. Admin Operations End to End", level=1)
    add_para(
        doc,
        "The admin console is reachable at /admin and uses staff credentials. Customer sessions are redirected away from the admin screen. Admin screens subscribe to order and catalog websocket events so dashboards, inventory, shipments, returns, customers, reports, audit logs, and stock alerts can refresh from live backend events.",
    )
    add_table(
        doc,
        ["Admin workflow", "Implementation", "Operational notes"],
        [
            ["Dashboard", "GET /api/admin/dashboard plus realtime refresh", "Shows KPIs and recent order activity."],
            ["Catalog/product upload", "AdminCatalogManager calls product/category/collection/variant/image endpoints", "Use quick-create for a product with initial variant; upload images through /api/admin/product-images."],
            ["Inventory", "GET/POST /api/admin/inventory", "Stock adjustments create ledger entries and publish catalog/inventory events."],
            ["Orders", "GET /api/admin/orders, PATCH status, POST workflow", "Use workflow actions for pack, label, pickup, transit, delivery, failure, RTO, refund-like paths."],
            ["Shipments", "GET/POST /api/admin/shipments, label/manifest downloads", "Manual courier is available; Shiprocket can be enabled with credentials."],
            ["Returns", "GET /api/admin/returns, PATCH /api/admin/returns/{id}", "Approve, reject, receive, refund according to return state."],
            ["Coupons", "GET/POST/PATCH /api/admin/coupons", "Coupons exist as first-class order model records and admin UI tab."],
            ["Customers/reports/audit", "Admin analytics endpoints", "Audit logs support traceability by action/entity/search query."],
        ],
        [1900, 3500, 3960],
    )
    doc.add_heading("8.1 Order Lifecycle Actions", level=2)
    add_table(
        doc,
        ["Lifecycle stage", "Customer-visible result", "Admin responsibility"],
        [
            ["Confirmed / Processing", "Order appears in My Orders and Tracking", "Review payment/COD and prepare stock."],
            ["Quality check / Packing", "Timeline reflects operation progress", "Inspect textile, blouse option, SKU, and package."],
            ["Label created / Pickup scheduled", "Tracking starts showing courier/AWB data", "Create shipment, add label, AWB, manifest, provider."],
            ["In transit / Out for delivery", "Realtime tracking and notifications update", "Apply courier webhook or manual event updates."],
            ["Delivered", "Invoice/review/return eligibility path opens", "Close fulfillment and verify payment/loyalty."],
            ["Delivery failed / RTO", "Customer sees failed/RTO timeline", "Update event note/location and decide reattempt/refund policy."],
            ["Cancelled / Returned / Refunded", "Order status and payment/refund state update", "Use cancel/return/refund flows and audit the action."],
        ],
        [2300, 3300, 3760],
    )

    doc.add_heading("9. API Contract", level=1)
    add_para(doc, "All public backend routes are under /api except /health, /readiness, Django admin, static media, and websocket routes.")
    add_table(
        doc,
        ["Domain", "Routes"],
        [
            ["Auth", "POST /api/auth/otp/send; POST /api/auth/otp/verify; POST /api/auth/admin/login; POST /api/auth/refresh; POST /api/auth/logout; GET/PATCH /api/auth/me"],
            ["Addresses", "GET/POST /api/addresses; PATCH/DELETE /api/addresses/{address_id}"],
            ["Catalog", "GET /api/categories; GET /api/collections; GET /api/catalog/facets; GET /api/products; GET /api/search; GET /api/products/{slug}; GET /api/products/{slug}/delivery"],
            ["Reviews", "GET/POST /api/products/{slug}/reviews"],
            ["Cart/wishlist", "GET/POST/DELETE /api/cart; PATCH/DELETE /api/cart/items/{id}; POST /api/cart/coupon; GET /api/checkout/summary; GET/POST /api/wishlist; DELETE /api/wishlist/{product_slug}"],
            ["Orders/returns", "GET/POST /api/orders; GET /api/orders/track; GET /api/orders/{id}; GET /api/orders/{id}/invoice; POST /api/orders/{id}/cancel; GET/POST /api/returns"],
            ["Payments", "POST /api/payments/razorpay/order; POST /api/payments/razorpay/verify; POST /api/payments/razorpay/webhook; POST /api/payments/webhook; POST /api/payments/refund"],
            ["Loyalty", "GET /api/loyalty/balance; GET /api/loyalty/history; GET /api/loyalty/rewards; POST /api/loyalty/redeem/{reward_id}"],
            ["Notifications", "GET /api/notifications?page=&per_page=; GET /api/notifications/count; PATCH /api/notifications"],
            ["AI", "POST /api/ai/tryon; POST /api/ai/voice-search; GET/POST /api/ai/recommend"],
            ["Admin catalog", "GET/POST /api/admin/products; POST /api/admin/products/quick-create; PATCH/DELETE /api/admin/products/{id}; GET/POST /api/admin/variants; PATCH /api/admin/variants/{id}; categories/collections/images"],
            ["Admin operations", "Dashboard, inventory, unsold alerts, orders, workflow, shipments, returns, customers, reports, audit logs, coupons, invoice, label, manifest endpoints under /api/admin/"],
            ["Health/docs", "GET /health; GET /api/health; GET /readiness; GET /api/readiness; GET /api/schema; GET /api/docs"],
        ],
        [1800, 7560],
    )
    add_callout(
        doc,
        "API documentation",
        "The OpenAPI schema is generated by drf-spectacular at /api/schema and rendered at /api/docs. Treat that page as the exact request/response reference during frontend or integration work.",
    )

    doc.add_heading("10. Realtime and Background Jobs", level=1)
    add_table(
        doc,
        ["Realtime channel", "Endpoint", "Purpose"],
        [
            ["Catalog", "/ws/catalog/", "Product, category, collection, image, variant, and inventory update events."],
            ["Notifications", "/ws/notifications/", "New notification and mark-read events with unread count."],
            ["Orders", "/ws/orders/", "Admin/global order updates."],
            ["Single order", "/ws/orders/{order_id}/", "Customer or admin tracking updates for one order."],
        ],
        [1800, 2500, 5060],
    )
    add_bullets(
        doc,
        [
            "WebSocket authentication uses JWT through the csm-token subprotocol.",
            "The frontend realtime client refreshes expired access tokens and reconnects after 4401 or heartbeat timeout.",
            "Heartbeat interval is 25 seconds and stale socket timeout is 60 seconds.",
            "Vite proxies /ws to Django during local development; production should route /ws to Daphne through Nginx.",
            "Redis channel layer is mandatory in production; in-memory channels are development only.",
        ],
    )
    doc.add_heading("10.1 Celery", level=2)
    add_para(
        doc,
        "Celery worker and beat are defined in docker-compose.yml and docker-compose.prod.yml. The beat schedule currently includes release-expired-stock-reservations every 300 seconds. Future async work should use Celery for email, WhatsApp, SMS, invoice generation, report generation, and slow external provider calls.",
    )

    doc.add_heading("11. Payments, Shipping, Invoices, Returns, Loyalty, and Notifications", level=1)
    add_table(
        doc,
        ["Subsystem", "Current behavior", "Production requirement"],
        [
            ["Pricing/GST", "Checkout totals include subtotal, discount, CGST, SGST, shipping, and total.", "Validate GST rate and HSN with accountant before launch."],
            ["Coupons", "Coupon model and admin UI/API support create/update and checkout application.", "Define business rules, limits, expiry, and audit reporting."],
            ["Loyalty", "Points ledger and rewards are exposed to customer account and checkout/order logic.", "Approve earn/redeem formula and expiry policy."],
            ["Razorpay", "Order creation, verification, webhooks, and refunds are implemented.", "Register webhook URL, set webhook secret, run test payment, switch to live keys only after UAT."],
            ["COD", "Supported as payment method without gateway dependency.", "Define COD availability, max order value, and courier policy."],
            ["Invoices", "Customer/admin invoice download endpoints exist.", "Verify GST invoice format, legal seller details, and PDF rendering on production host."],
            ["Shipping", "Manual courier label/manifest and Shiprocket adapter placeholders are present.", "Use manual provider or configure Shiprocket credentials/webhooks."],
            ["Returns/RTO", "Return requests and admin return status updates exist.", "Publish return window, eligibility, refund timelines, and damaged-item process."],
            ["Notifications", "App feed is realtime and paginated; Resend/Twilio/Gupshup integration settings exist.", "Enable at least one customer notification channel and test order status templates."],
            ["AI try-on", "Claude-ready endpoint with model/session storage.", "Set ANTHROPIC_API_KEY for real vision-based try-on; keep non-critical to checkout."],
        ],
        [1700, 4200, 3460],
    )

    doc.add_heading("12. Security and Compliance Controls", level=1)
    add_bullets(
        doc,
        [
            "JWT access/refresh tokens are used for API authentication; frontend refreshes access tokens before requests.",
            "Admin APIs use staff/superuser checks and are separated from customer routes.",
            "Production deploy checks block unsafe fallbacks, missing live OTP channels, missing notification channels, missing payment secrets, SQLite, and in-memory realtime.",
            "Razorpay signatures and webhook signatures are verified server side; duplicate webhooks are tracked through RazorpayWebhookEvent.",
            "Courier webhooks have signature helpers; Shiprocket requires credentials and webhook secret if selected.",
            "Throttling is configured for anonymous, user, OTP, admin login, tracking, and courier webhook traffic.",
            "Production security defaults include HTTPS redirect, secure cookies, HSTS, and forwarded-proto support when DEBUG=False.",
            "Never expose real provider keys in documentation, screenshots, commits, issue comments, or chat. Rotate any credential that was pasted into a conversation before production use.",
        ],
    )
    add_table(
        doc,
        ["Risk", "Control", "Operator action"],
        [
            ["OTP abuse", "DRF scoped throttles and production live channel requirement", "Tune DRF_OTP_THROTTLE and monitor provider usage."],
            ["Payment spoofing", "Razorpay HMAC verification and no dev fallback in production", "Keep PAYMENT_DEV_FALLBACK_ENABLED=False and verify webhook secret."],
            ["Overselling", "Variant stock, stock reservations, ledger, and reservation expiry task", "Run Celery beat and test last-SKU concurrency before launch."],
            ["Stale realtime", "Redis channel layer and frontend heartbeat/reconnect", "Use Redis in production and route /ws through Nginx to Daphne."],
            ["Data loss", "PostgreSQL persistent volume", "Implement external scheduled backups and restore drills."],
            ["Secret leak", ".env ignored and .env.example uses placeholders", "Rotate any exposed keys and restrict GitHub/environment access."],
        ],
        [1800, 3800, 3760],
    )

    doc.add_heading("13. Deployment and CI/CD", level=1)
    doc.add_heading("13.1 GitHub CI", level=2)
    add_para(
        doc,
        "GitHub Actions runs on pushes and pull requests to main. The backend job uses Python 3.12, PostgreSQL 16, installs requirements and system libraries for PDF support, runs Django checks, migrations, and the backend test suite. The frontend job uses Node 22, npm ci, ESLint, and production build.",
    )
    add_code_block(
        doc,
        [
            "python backend/manage.py check",
            "python backend/manage.py migrate --noinput",
            "python backend/manage.py test accounts catalog cart orders payments inventory loyalty notifications analytics shipping reviews ai",
            "cd frontend && npm run lint",
            "cd frontend && npm run build",
        ],
    )
    doc.add_heading("13.2 Production Docker Compose", level=2)
    add_table(
        doc,
        ["Service", "Purpose", "Production notes"],
        [
            ["nginx", "Public HTTP/HTTPS reverse proxy", "Routes frontend, API, and /ws traffic; mount SSL certificate paths."],
            ["api", "Django ASGI app via Daphne", "Runs check --deploy, migrate, then daphne. Must pass readiness checks."],
            ["frontend", "Builds and serves React dist", "Uses Node 22 alpine and serve. In larger deployments, replace with CDN/static hosting."],
            ["celery_worker", "Background jobs", "Run with Redis/PostgreSQL and monitor queue health."],
            ["celery_beat", "Scheduled jobs", "Required for expired reservation release."],
            ["db", "PostgreSQL 16", "Use managed PostgreSQL or persistent volume plus backups."],
            ["redis", "Broker and channel layer", "Use managed Redis for production if possible."],
        ],
        [1700, 3000, 4660],
    )
    doc.add_heading("13.3 Deployment Checklist", level=2)
    add_numbers(
        doc,
        [
            "Provision production PostgreSQL and Redis.",
            "Create .env with production values; never reuse dev SECRET_KEY or default passwords.",
            "Set APP_ENV=production and DEBUG=False.",
            "Configure domain, ALLOWED_HOSTS, CORS_ALLOWED_ORIGINS, CSRF_TRUSTED_ORIGINS, TLS certificates, and Nginx routing.",
            "Configure Razorpay keys and webhook secret; register payment webhook URL.",
            "Configure OTP provider: Twilio SMS or Resend email OTP.",
            "Configure customer notifications: Resend email and/or Gupshup WhatsApp.",
            "Choose DEFAULT_COURIER_PROVIDER=manual or configure Shiprocket credentials and webhook secret.",
            "Run python backend/manage.py check --deploy.",
            "Run migrations and seed only the data intended for production.",
            "Smoke test login, catalog, cart, checkout, payment test, COD order, invoice, admin workflow, tracking, notification, return.",
            "Enable backups, log retention, uptime checks, and provider webhook monitoring.",
        ],
    )

    doc.add_heading("14. Operational Runbooks", level=1)
    doc.add_heading("14.1 Daily Admin Checklist", level=2)
    add_bullets(
        doc,
        [
            "Open /admin and confirm dashboard KPIs load.",
            "Review new orders and payment status.",
            "Move paid/COD-confirmed orders to quality check and packing.",
            "Create or update shipments with provider, AWB, label, manifest, location, and customer note.",
            "Monitor returns and delivery failed/RTO states.",
            "Review low/unsold inventory alerts.",
            "Check notification delivery failures if providers are enabled.",
            "Export or review daily sales/refund/return reports.",
        ],
    )
    doc.add_heading("14.2 Incident Runbook", level=2)
    add_table(
        doc,
        ["Incident", "First checks", "Recovery action"],
        [
            ["Frontend unavailable", "Nginx/frontend container logs, DNS, TLS, / health", "Restart frontend/Nginx, rollback build, verify static dist."],
            ["API unavailable", "/api/health, /api/readiness, API logs, DB/Redis health", "Restart API, verify migrations, DB connection, environment variables."],
            ["Login OTP failing", "Readiness OTP channel, Twilio/Resend logs, DRF throttle", "Fix provider credentials, disable bad channel, communicate fallback support process."],
            ["Payments failing", "Razorpay dashboard, payment verify/webhook logs, webhook secret", "Pause prepaid checkout if needed, use COD, reconcile pending payments."],
            ["Realtime not updating", "Redis, Daphne, Nginx /ws routing, browser console", "Restart Redis/API, verify CHANNEL_LAYER_BACKEND=redis and /ws proxy."],
            ["Oversell/stock mismatch", "StockLedger, StockReservation, order items, Celery beat", "Pause SKU, adjust inventory with audit note, reconcile open reservations."],
            ["Shipping webhook failing", "Courier webhook logs and signature settings", "Switch to manual event updates until provider webhook is restored."],
        ],
        [1800, 3600, 3960],
    )
    doc.add_heading("14.3 Backup and Restore", level=2)
    add_bullets(
        doc,
        [
            "Back up PostgreSQL at least daily before launch; increase frequency as order volume grows.",
            "Back up uploaded product images/media if stored locally; prefer object storage for production media.",
            "Test restore into a staging environment monthly.",
            "Keep .env and provider secrets in a secure secrets manager, not only on one server.",
            "Document RPO/RTO with the business owner before taking real orders.",
        ],
    )

    doc.add_heading("15. Testing, QA, and Acceptance Checklist", level=1)
    add_table(
        doc,
        ["Test area", "Command or flow", "Expected result"],
        [
            ["Backend system check", "python backend/manage.py check", "No system check errors."],
            ["Backend test suite", "python backend/manage.py test accounts catalog cart orders payments inventory loyalty notifications analytics shipping reviews ai", "All app tests pass."],
            ["Production deploy gate", "DEBUG=False APP_ENV=production python backend/manage.py check --deploy", "Blocks missing production requirements; passes only with proper env."],
            ["Frontend lint", "cd frontend && npm run lint", "No ESLint errors."],
            ["Frontend build", "cd frontend && npm run build", "TypeScript and Vite build pass."],
            ["Customer smoke", "Login/signup -> browse -> product -> cart -> checkout -> COD/Razorpay test -> order detail", "Order exists with correct totals, stock, invoice, notifications."],
            ["Admin smoke", "Admin login -> add product/image/variant -> stock -> publish -> customer sees product", "Catalog updates appear and realtime refresh works."],
            ["Tracking smoke", "Admin updates shipment -> customer tracking/order page", "Timeline and websocket status update."],
            ["Return smoke", "Delivered order -> return request -> admin decision -> refund status", "Return lifecycle and notifications/audit logs work."],
        ],
        [1900, 4300, 3160],
    )
    doc.add_heading("15.1 Launch Acceptance Criteria", level=2)
    add_bullets(
        doc,
        [
            "A real seeded or production customer can authenticate through a live OTP channel.",
            "A real SKU can be added to cart, reserved, ordered, paid or marked COD, invoiced, shipped, tracked, delivered, reviewed, and returned if eligible.",
            "Admin can upload a new collection/product/image/variant, adjust stock, publish it, and see it on the customer storefront.",
            "Razorpay webhook and verify flows are tested with duplicate and invalid signatures.",
            "Customer order updates reach at least one enabled notification channel.",
            "Production readiness endpoint reports ready, not degraded.",
            "No production route depends on mock/local cart/order/admin data.",
            "Backups, SSL, domain, logs, and rollback procedure are ready.",
        ],
    )

    doc.add_heading("16. Known Limits and Next Roadmap", level=1)
    add_table(
        doc,
        ["Area", "Current status", "Recommended next action"],
        [
            ["Multi-vendor marketplace", "Out of V1 scope", "Keep single-brand until inventory/order ops are stable."],
            ["Native mobile apps", "Out of V1 scope", "Use responsive web/PWA first."],
            ["Warehouse scanning", "Not implemented", "Add barcode/QR receiving and packing after order volume grows."],
            ["Advanced search/ML", "Basic API filtering/search plus AI helper", "Add Meilisearch/Typesense/Elastic only after catalog size demands it."],
            ["AI try-on", "Claude-ready non-critical enhancement", "Enable only after ANTHROPIC_API_KEY and privacy policy are production-ready."],
            ["Provider webhooks", "Razorpay and courier webhook paths exist", "Complete live provider registration and staging webhook tests."],
            ["Observability", "Health/readiness and logs exist", "Add Sentry, structured logs, metrics, uptime checks, and alerts."],
            ["Media storage", "Local media in dev", "Move production media to object storage/CDN."],
        ],
        [1900, 3300, 4160],
    )
    add_callout(
        doc,
        "Final go-live advice",
        "Do not take real payments until production readiness is green, payment/OTP/notification providers are live, webhook replay has been tested, and at least one full order has been traced from customer checkout through admin shipment to delivered status.",
    )

    doc.add_heading("Appendix A. Command Reference", level=1)
    add_table(
        doc,
        ["Action", "Command"],
        [
            ["Seed local data", "python backend/manage.py seed_csm"],
            ["Create migrations", "python backend/manage.py makemigrations"],
            ["Apply migrations", "python backend/manage.py migrate"],
            ["Backend check", "python backend/manage.py check"],
            ["Backend tests", "python backend/manage.py test accounts catalog cart orders payments inventory loyalty notifications analytics shipping reviews ai"],
            ["Run ASGI backend", "cd backend && daphne -b 0.0.0.0 -p 8000 csm_backend.asgi:application"],
            ["Frontend install", "cd frontend && npm install"],
            ["Frontend dev", "cd frontend && npm run dev"],
            ["Frontend lint", "cd frontend && npm run lint"],
            ["Frontend build", "cd frontend && npm run build"],
            ["Docker dev", "docker compose up"],
            ["Docker prod", "docker compose -f docker-compose.prod.yml up"],
        ],
        [2600, 6760],
    )

    doc.add_heading("Appendix B. Source Files to Know", level=1)
    add_table(
        doc,
        ["File", "Why it matters"],
        [
            ["backend/csm_backend/settings.py", "Environment, security, database, REST, JWT, Redis, Celery, provider settings."],
            ["backend/csm_backend/urls.py", "Root HTTP API routing and health/docs endpoints."],
            ["backend/csm_backend/asgi.py + routing.py", "ASGI and websocket routing."],
            ["backend/accounts/checks.py", "Production deploy safety checks."],
            ["backend/orders/services.py", "Checkout, stock, order lifecycle, loyalty/invoice side effects."],
            ["backend/orders/views.py", "Customer/admin order, return, coupon, invoice endpoints."],
            ["backend/payments/services.py", "Razorpay signature, webhook, order creation, refund support."],
            ["backend/shipping/services.py", "Manual/Shiprocket shipment and tracking event behavior."],
            ["backend/notifications/services.py", "Notification creation and provider sends."],
            ["frontend/src/lib/api.ts", "Single frontend API contract and token refresh behavior."],
            ["frontend/src/lib/realtime.ts", "WebSocket connection, heartbeat, reconnect, token refresh."],
            ["frontend/src/App.tsx", "Route protection and customer/admin route map."],
            ["frontend/src/pages/Admin.tsx", "Admin operational console."],
            ["frontend/src/pages/Checkout.tsx", "Customer checkout and payment flow."],
            ["docker-compose.prod.yml", "Production runtime wiring and deploy check gate."],
            [".github/workflows/ci.yml", "Backend and frontend CI gates."],
        ],
        [3100, 6260],
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
